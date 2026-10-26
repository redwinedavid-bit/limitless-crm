"""Convert the read-only active Job Board Excel export into migration CSVs."""
import csv,sys
from pathlib import Path
from openpyxl import load_workbook
src=Path(sys.argv[1]); out=Path(sys.argv[2]);out.mkdir(parents=True,exist_ok=True)
wb=load_workbook(src,read_only=True,data_only=True)
name_map={'Jobs OpenCosting':'Jobs Open_Costing.csv','Customer':'Customer.csv','David POs':'David POs.csv','Adrian POs':'Adrian POs.csv','Danny POs':'Danny POs.csv','David Hrs Mileage':'David Hrs Mileage.csv','Adrian Hrs Mileage':'Adrian Hrs Mileage.csv','Danny Hrs Mileage':'Danny Hrs Mileage.csv'}
for name,filename in name_map.items():
 ws=wb[name]; rows=ws.iter_rows(values_only=True)
 head=next(rows);heads=[str(x).strip() if x is not None else f'__unnamed_{i}' for i,x in enumerate(head)]
 if name=='Jobs OpenCosting': heads=[h if h!='Customer' else 'Customer ' for h in heads]
 with (out/filename).open('w',newline='',encoding='utf-8') as f:
  w=csv.writer(f);w.writerow(heads)
  for row in rows:
   if not any(x is not None and str(x).strip() for x in row):continue
   w.writerow(['' if x is None else x.isoformat() if hasattr(x,'isoformat') else str(x) for x in row])
 print(name,filename)
